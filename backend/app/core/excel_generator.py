import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows


def generate_excel_from_blueprint(blueprint: dict, title: str) -> io.BytesIO:
    """Takes an AI blueprint JSON and dynamically generates a rich, native Excel
    dashboard with charts using openpyxl, saving it to a BytesIO buffer.
    """
    output = io.BytesIO()
    wb = Workbook()

    # Check if this is a BI Dashboard (contains "charts")
    if "charts" in blueprint and isinstance(blueprint["charts"], list):
        ws = wb.active
        ws.title = "Dashboard Summary"

        # Add Dashboard Title and Summary
        ws.merge_cells("A1:F1")
        title_cell = ws["A1"]
        title_cell.value = blueprint.get("dashboardTitle", title)
        title_cell.font = Font(size=20, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(
            start_color="10B981", end_color="10B981", fill_type="solid",
        )  # Emerald Green
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        ws.merge_cells("A3:F4")
        summary_cell = ws["A3"]
        summary_cell.value = blueprint.get("dashboardSummary", "")
        summary_cell.font = Font(size=12, italic=True)
        summary_cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Generate a separate sheet for each chart
        charts_data = blueprint["charts"]
        for idx, chart in enumerate(charts_data):
            sheet_title = (chart.get("title") or f"Chart_{idx}")[:30].replace("/", "-")
            c_ws = wb.create_sheet(title=sheet_title)

            c_ws.merge_cells("A1:C1")
            c_ws["A1"] = chart.get("title", "Data")
            c_ws["A1"].font = Font(size=14, bold=True)

            c_ws.merge_cells("A2:C2")
            c_ws["A2"] = chart.get("description", "")
            c_ws["A2"].alignment = Alignment(wrap_text=True)
            c_ws.row_dimensions[2].height = 30

            data = chart.get("data", [])
            if not data:
                continue

            # Put data into sheet starting at row 4
            df = pd.DataFrame(data)
            # Make sure df has 'name' and 'value' cols
            if "name" not in df.columns or "value" not in df.columns:
                # If they used different keys, just take the first two
                cols = list(df.columns)
                if len(cols) >= 2:
                    df = df[[cols[0], cols[1]]]
            else:
                df = df[["name", "value"]]

            # Rename headers
            df.columns = [
                chart.get("xAxisKey", "Category"),
                chart.get("yAxisKey", "Value"),
            ]

            for r_idx, row in enumerate(
                dataframe_to_rows(df, index=False, header=True), start=4,
            ):
                for c_idx, value in enumerate(row, start=1):
                    cell = c_ws.cell(row=r_idx, column=c_idx, value=value)
                    if r_idx == 4:  # Header
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(
                            start_color="E5E7EB", end_color="E5E7EB", fill_type="solid",
                        )

            # Format column widths
            c_ws.column_dimensions["A"].width = 25
            c_ws.column_dimensions["B"].width = 15

            # Create Native Excel Chart
            chart_type = chart.get("chartType", "bar").lower()
            max_row = 4 + len(data)

            x_values = Reference(c_ws, min_col=1, min_row=5, max_row=max_row)
            y_values = Reference(
                c_ws, min_col=2, min_row=4, max_row=max_row,
            )  # includes header

            if chart_type == "bar":
                xl_chart = BarChart()
                xl_chart.type = "col"
                xl_chart.style = 27
            elif chart_type == "line":
                xl_chart = LineChart()
                xl_chart.style = 27
            elif chart_type == "pie":
                xl_chart = PieChart()
            elif chart_type == "area":
                xl_chart = AreaChart()
            else:
                xl_chart = BarChart()  # Fallback

            xl_chart.title = chart.get("title", "")
            if chart_type != "pie":
                xl_chart.y_axis.title = chart.get("yAxisKey", "Value")
                xl_chart.x_axis.title = chart.get("xAxisKey", "Category")

            xl_chart.add_data(y_values, titles_from_data=True)
            xl_chart.set_categories(x_values)

            # Add chart to sheet
            c_ws.add_chart(xl_chart, "E4")

    elif "predictedTrendline" in blueprint and isinstance(
        blueprint["predictedTrendline"], list,
    ):
        # ── ML Trend Forecast Report ──
        ws = wb.active
        ws.title = "Forecast Summary"

        # Title
        ws.merge_cells("A1:E1")
        title_cell = ws["A1"]
        title_cell.value = blueprint.get("forecastTitle", title)
        title_cell.font = Font(size=18, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(
            start_color="059669", end_color="059669", fill_type="solid",
        )  # Emerald
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # Key Metrics Row
        ws["A3"] = "Projected Growth"
        ws["A3"].font = Font(bold=True, color="4B5563")
        ws["B3"] = blueprint.get("projectedGrowth", blueprint.get("growthRate", "N/A"))
        ws["B3"].font = Font(size=14, bold=True, color="059669")

        ws["C3"] = "Trend Direction"
        ws["C3"].font = Font(bold=True, color="4B5563")
        ws["D3"] = blueprint.get("trendDirection", "Upward")
        ws["D3"].font = Font(size=14, bold=True, color="2563EB")

        # Executive Summary
        ws.merge_cells("A5:E6")
        summary_cell = ws["A5"]
        summary_cell.value = blueprint.get("executiveSummary", "")
        summary_cell.font = Font(size=11, italic=True)
        summary_cell.alignment = Alignment(vertical="top", wrap_text=True)

        # Create Forecast Data Sheet with Native Line Chart
        c_ws = wb.create_sheet(title="Predictive Trendline")

        headers = [
            "Period",
            "Historical Value",
            "Predicted Value",
            "Pessimistic Bound (P10)",
            "Optimistic Bound (P90)",
        ]
        for c_idx, h in enumerate(headers, start=1):
            cell = c_ws.cell(row=3, column=c_idx, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="1F2937", end_color="1F2937", fill_type="solid",
            )
            cell.alignment = Alignment(horizontal="center")
        c_ws.row_dimensions[3].height = 25

        trend_data = blueprint["predictedTrendline"]
        for r_idx, pt in enumerate(trend_data, start=4):
            c_ws.cell(row=r_idx, column=1, value=pt.get("period"))
            c_ws.cell(row=r_idx, column=2, value=pt.get("historicalValue"))
            c_ws.cell(row=r_idx, column=3, value=pt.get("predictedValue"))
            c_ws.cell(row=r_idx, column=4, value=pt.get("pessimisticBound"))
            c_ws.cell(row=r_idx, column=5, value=pt.get("optimisticBound"))

        c_ws.column_dimensions["A"].width = 20
        c_ws.column_dimensions["B"].width = 18
        c_ws.column_dimensions["C"].width = 18
        c_ws.column_dimensions["D"].width = 24
        c_ws.column_dimensions["E"].width = 24

        # Add Native Excel LineChart
        if trend_data:
            chart = LineChart()
            chart.title = blueprint.get("forecastTitle", "Predictive Forecast Model")
            chart.style = 13
            chart.y_axis.title = blueprint.get("targetMetric", "Value")
            chart.x_axis.title = "Period"
            chart.width = 18
            chart.height = 12

            data_ref = Reference(
                c_ws, min_col=2, min_row=3, max_col=5, max_row=3 + len(trend_data),
            )
            cats_ref = Reference(
                c_ws, min_col=1, min_row=4, max_row=3 + len(trend_data),
            )
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            c_ws.add_chart(chart, "G3")

    elif "missingValues" in blueprint or "aiInsights" in blueprint:
        # This is an Executive Summary / AI Analysis
        ws = wb.active
        ws.title = "Executive Summary"

        ws.merge_cells("A1:B1")
        title_cell = ws["A1"]
        title_cell.value = title
        title_cell.font = Font(size=18, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(
            start_color="10B981", end_color="10B981", fill_type="solid",
        )
        ws.row_dimensions[1].height = 30

        row_num = 3
        for key, value in blueprint.items():
            if isinstance(value, dict):
                ws.cell(row=row_num, column=1, value=str(key).title()).font = Font(
                    bold=True,
                )
                row_num += 1
                for k, v in value.items():
                    ws.cell(row=row_num, column=1, value=str(k))
                    ws.cell(row=row_num, column=2, value=str(v))
                    row_num += 1
            elif isinstance(value, list):
                ws.cell(row=row_num, column=1, value=str(key).title()).font = Font(
                    bold=True,
                )
                row_num += 1
                for item in value:
                    ws.cell(row=row_num, column=2, value=str(item))
                    row_num += 1
            else:
                ws.cell(row=row_num, column=1, value=str(key).title()).font = Font(
                    bold=True,
                )
                ws.cell(row=row_num, column=2, value=str(value)).alignment = Alignment(
                    wrap_text=True,
                )
                row_num += 1
            row_num += 1

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 80

    else:
        # Generic fallback
        ws = wb.active
        ws.title = "Report Data"
        df_data = []
        for key, value in blueprint.items():
            if isinstance(value, (dict, list)):
                value = str(value)
            df_data.append({"Section": key, "Content": value})

        df = pd.DataFrame(df_data)
        for r_idx, row in enumerate(
            dataframe_to_rows(df, index=False, header=True), start=1,
        ):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 80

    wb.save(output)
    output.seek(0)
    return output
